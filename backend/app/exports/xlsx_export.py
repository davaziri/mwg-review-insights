"""Excel workbook: About, Theme summary, Coded reviews and Monthly trends sheets.

Kept to what Google Sheets opens cleanly: plain values (no formulas at all), bold headers, frozen
header rows, column widths, percent and date formats, and links.
"""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.exports.common import Brand
from app.exports.tables import ExportContext, about, coded_reviews, monthly_trends, theme_summary

MAX_CELL = 32_767  # the most characters a spreadsheet cell holds


def _put(ws, row: int, col: int, value, *, percent: bool = False):
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if len(value) > MAX_CELL:
            value = value[: MAX_CELL - 20] + " [cut: too long]"
    cell.value = value
    if isinstance(value, str):
        # Always text: a review that starts with "=" must never become a formula.
        cell.data_type = "s"
    elif isinstance(value, date):
        cell.number_format = "yyyy-mm-dd"
    elif percent and isinstance(value, float):
        cell.number_format = "0%"
    return cell


def _table(ws, header: list[str], rows: list[list], brand: Brand, *, percent_cols: set[int] = frozenset(), widths: dict[int, int] | None = None, link_col: int | None = None):
    head_fill = PatternFill("solid", fgColor=brand.palette["soft"].lstrip("#"))
    for c, name in enumerate(header, start=1):
        cell = _put(ws, 1, c, name)
        cell.font = Font(bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = _put(ws, r, c, value, percent=c in percent_cols)
            if link_col == c and isinstance(value, str) and value.startswith("https://"):
                cell.hyperlink = value
                cell.font = Font(color=brand.palette["link"].lstrip("#"), underline="single")
    ws.freeze_panes = "A2"
    for c in range(1, len(header) + 1):
        ws.column_dimensions[get_column_letter(c)].width = (widths or {}).get(c, max(10, min(40, len(header[c - 1]) + 4)))


def workbook(session: Session, ctx: ExportContext, brand: Brand) -> bytes:
    wb = Workbook()

    ws = wb.active
    ws.title = "About"
    ws.cell(row=1, column=1, value=f"{brand.company_name} {brand.app_name}").font = Font(bold=True, size=14)
    for r, (label, value) in enumerate(about(ctx), start=3):
        _put(ws, r, 1, label).font = Font(bold=True)
        _put(ws, r, 2, value).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 100

    header, rows = theme_summary(ctx)
    _table(wb.create_sheet("Theme summary"), header, rows, brand, percent_cols={4}, widths={1: 30, 2: 70})

    header, rows = coded_reviews(session, ctx)
    n = len(header)
    _table(
        wb.create_sheet("Coded reviews"),
        header,
        rows,
        brand,
        widths={1: 14, 2: 22, 3: 12, 4: 16, 6: 12, n - 1: 50, n: 80},
        link_col=2,
    )

    header, rows = monthly_trends(ctx)
    _table(wb.create_sheet("Monthly trends"), header, rows, brand, percent_cols={4, *range(9, len(header) + 1)}, widths={1: 10})

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
