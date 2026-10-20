"""Exports: CSV and Excel of coded reviews, and the PDF reports, all respecting the filters."""

import csv
import io
import os
import re
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.analysis.aggregate import Filters
from app.exports.common import describe_filters, export_name, load_brand
from app.exports.portfolio_report import portfolio_report_html
from app.exports.report import title_report_html
from app.exports.tables import build_context
from app.models import Review, Title
from app.steam.sync import upsert_reviews
from app.timeutil import utcnow
from tests.fakes import make_review
from tests.test_portfolio_api import analyze, seed

APP = 3308200
TRICKY = "=HYPERLINK(\"http://example.com\",\"click\") but the combat is fun."


@pytest.fixture
def analyzed(api, session, session_factory, settings, fake_llm):
    seed(session, 1, APP)
    upsert_reviews(session, 1, APP, [make_review(99_999, 1_765_000_000, text=TRICKY)], seen_at=utcnow())
    session.commit()
    analyze(api, session_factory, settings, fake_llm, 1)
    return api


def test_export_names():
    today = date(2026, 10, 20)
    assert export_name("Duck Paradox", "Report", "pdf", today) == "MWG_Duck_Paradox_Report_2026-10-20.pdf"
    assert export_name("Floor is...What!?", "Data", "xlsx", today) == "MWG_Floor_is_What_Data_2026-10-20.xlsx"
    assert export_name("Hive Jump 2: Survivors", "CodedReviews", "csv", today) == "MWG_Hive_Jump_2_Survivors_CodedReviews_2026-10-20.csv"


def test_filters_described_in_words():
    f = Filters(date_from=date(2026, 1, 1), voted_up=False, playtime_bands=["lt2h"], languages=["schinese"], steam_deck=True)
    assert describe_filters(f) == [
        "Reviews written from 2026-01-01 to today",
        "Not recommended reviews only",
        "Playtime: Under 2h",
        "Languages: Chinese (Simplified)",
        "Played mostly on Steam Deck",
    ]
    assert describe_filters(Filters()) == []


def test_csv_of_coded_reviews(analyzed):
    resp = analyzed.get("/api/titles/1/export.csv")
    assert resp.status_code == 200
    assert re.search(r'filename="MWG_Tombwater_CodedReviews_\d{4}-\d{2}-\d{2}\.csv"', resp.headers["content-disposition"])
    assert resp.content.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))
    header, body = rows[0], rows[1:]
    assert len(body) == 64  # every review, analyzed or not
    assert {"Combat feel", "Price and value", "Bugs and crashes", "Art style", "Overall impression", "Review text"} <= set(header)
    col = {name: i for i, name in enumerate(header)}
    analyzed_rows = [r for r in body if r[col["Analyzed"]] == "Yes"]
    assert len(analyzed_rows) == 55 and {r[col["Analyzed"]] for r in body} == {"Yes", "No, too short"}
    first = next(r for r in body if r[col["Review ID"]] == "10000")
    assert first[col["Combat feel"]] == "Positive" and first[col["Bugs and crashes"]] == ""
    assert first[col["Overall impression quote"]] == "Great game, I recommend it."
    assert first[col["Steam link"]].startswith("https://steamcommunity.com/profiles/")
    # A review that looks like a formula is kept as plain text.
    tricky = next(r for r in body if r[col["Review ID"]] == "99999")
    assert tricky[col["Review text"]] == "'" + TRICKY
    # Filters apply.
    negative = list(csv.reader(io.StringIO(analyzed.get("/api/titles/1/export.csv", params={"vote": "negative"}).content.decode("utf-8-sig"))))
    assert len(negative) - 1 == 18


def test_excel_workbook(analyzed):
    resp = analyzed.get("/api/titles/1/export.xlsx", params={"language": ["english"]})
    assert resp.status_code == 200 and "MWG_Tombwater_Data_" in resp.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == ["About", "Theme summary", "Coded reviews", "Monthly trends"]
    cells = [c for ws in wb.worksheets for row in ws.iter_rows() for c in row]
    assert not [c for c in cells if c.data_type == "f"]  # no formulas at all
    about = {ws_row[0].value: ws_row[1].value for ws_row in wb["About"].iter_rows(min_row=3) if ws_row[0].value}
    assert about["Filters"] == "Languages: English" and about["Codebook"].startswith("Version 1")
    reviews = wb["Coded reviews"]
    header = [c.value for c in reviews[1]]
    assert reviews.max_row - 1 == 55  # 64 reviews minus 9 in Chinese
    text_col = header.index("Review text") + 1
    tricky = next(r for r in reviews.iter_rows(min_row=2) if r[0].value == "99999")
    assert tricky[text_col - 1].value == TRICKY and tricky[text_col - 1].data_type == "s"
    assert tricky[1].hyperlink is not None
    summary = wb["Theme summary"]
    assert summary["D2"].number_format == "0%" and 0 < summary["D2"].value <= 1
    months = wb["Monthly trends"]
    assert [c.value for c in months[1]][:4] == ["Month", "Reviews", "Recommended", "Recommended share"]
    assert sum(r[1].value for r in months.iter_rows(min_row=2)) == 55


def test_title_report_html(analyzed, session, settings):
    title = session.get(Title, 1)
    ctx = build_context(session, settings, title, APP, Filters())
    html = title_report_html(session, ctx, load_brand(settings))
    for heading in ("Steam review report", "At a glance", "Reviews per month", "Overall impression", "Executive summary", "Themes",
                    "What players praise most", "What players criticize most", "How this report was made"):
        assert heading in html
    assert "—" not in html  # no em dashes anywhere in the report
    # Every quote is the reviewer's own text and links to the review.
    texts = [r.display_text for r in session.query(Review).all()]
    quotes = re.findall(r'<div class="text">&ldquo;(.*?)&rdquo;</div>', html)
    assert quotes
    import html as html_lib

    for q in quotes:
        assert any(html_lib.unescape(q) in t for t in texts)
    assert html.count("Read on Steam") == len(quotes)
    assert "Model" in html and settings.model_tagging in html
    # Review text is escaped, never markup.
    assert "<script" not in html


def test_portfolio_report_html(analyzed, session, settings):
    html = portfolio_report_html(session, settings, load_brand(settings), None, None)
    assert "Portfolio report" in html and "Tombwater" in html and "No reviews yet" in html
    assert html.count('<section class="page">') == 1  # only Tombwater has reviews here
    assert "—" not in html


def _chromium(settings) -> str | None:
    """A Chromium to print with: CHROMIUM_PATH, a preinstalled one, or the one `playwright install` added."""
    if os.environ.get("CHROMIUM_PATH"):
        return os.environ["CHROMIUM_PATH"]
    if Path("/opt/pw-browsers/chromium").exists():
        return "/opt/pw-browsers/chromium"
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        path = p.chromium.executable_path
    return path if Path(path).exists() else None


def test_pdf_downloads(analyzed, settings):
    path = _chromium(settings)
    if path is None:
        pytest.skip("no Chromium available to draw PDFs")
    settings.chromium_path = path
    resp = analyzed.get("/api/titles/1/export.pdf", params={"vote": "positive"})
    assert resp.status_code == 200, resp.text
    assert resp.content.startswith(b"%PDF") and "MWG_Tombwater_Report_" in resp.headers["content-disposition"]
    assert resp.content.count(b"/Type /Page\n") + resp.content.count(b"/Type /Page ") + resp.content.count(b"/Type/Page") >= 3 or len(resp.content) > 20_000
    portfolio = analyzed.get("/api/portfolio/export.pdf", params={"date_from": "2025-01-01"})
    assert portfolio.status_code == 200 and portfolio.content.startswith(b"%PDF")
    assert "MWG_Portfolio_Report_" in portfolio.headers["content-disposition"]


def test_pdf_failure_is_a_clear_error(analyzed, settings):
    settings.chromium_path = "/nonexistent/chrome"
    resp = analyzed.get("/api/titles/1/export.pdf")
    assert resp.status_code == 503 and "Could not make the PDF" in resp.json()["detail"]
